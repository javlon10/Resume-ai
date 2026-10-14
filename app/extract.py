"""Resume fayllaridan matn ajratish: PDF, DOCX, TXT/MD."""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".text"}


class ExtractError(RuntimeError):
    pass


def _clean(text: str) -> str:
    """Ortiqcha bo'shliq va ko'rinmas belgilarni tozalaydi (token tejaydi)."""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace(" ", " ").replace("﻿", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _from_pdf(path: Path) -> str:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
            # Jadval ko'rinishidagi resumelar uchun: matn bo'sh qolsa jadvaldan ol
            if not pages[-1].strip():
                for tbl in page.extract_tables() or []:
                    rows = [
                        " | ".join(c for c in row if c)
                        for row in tbl
                        if any(row)
                    ]
                    if rows:
                        pages[-1] = "\n".join(rows)
    text = "\n\n".join(pages)
    if not text.strip():
        raise ExtractError(
            f"{path.name}: PDF'dan matn chiqmadi — ehtimol skan (rasm) PDF. "
            "OCR kerak yoki matnli versiyasini bering."
        )
    return text


def _from_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for tbl in d.tables:                    # ko'p resume shablonlari jadvalga asoslangan
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(path: str | Path) -> str:
    """Fayl turiga qarab matn qaytaradi."""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Resume fayli topilmadi: {p}")
    suf = p.suffix.lower()
    if suf == ".pdf":
        text = _from_pdf(p)
    elif suf == ".docx":
        text = _from_docx(p)
    elif suf in {".txt", ".md", ".text"}:
        text = p.read_text(encoding="utf-8", errors="replace")
    elif suf == ".doc":
        raise ExtractError(
            f"{p.name}: eski .doc format qo'llanmaydi. .docx yoki .pdf ga o'tkazing "
            "(masalan: libreoffice --headless --convert-to docx)."
        )
    else:
        raise ExtractError(f"{p.name}: qo'llanmaydigan format '{suf}'. Mumkin: {sorted(SUPPORTED)}")

    text = _clean(text)
    if len(text) < 50:
        raise ExtractError(f"{p.name}: matn juda qisqa ({len(text)} belgi) — fayl buzuq bo'lishi mumkin")
    return text


def find_resumes(target: str | Path) -> list[Path]:
    """Fayl yoki katalog berilsa, qo'llanadigan resume fayllari ro'yxatini qaytaradi."""
    p = Path(target)
    if p.is_file():
        return [p]
    if p.is_dir():
        return sorted(
            f for f in p.rglob("*")
            if f.is_file() and f.suffix.lower() in SUPPORTED
        )
    raise FileNotFoundError(f"Yo'l topilmadi: {p}")


def truncate(text: str, max_chars: int) -> tuple[str, bool]:
    """Kontekstga sig'ishi uchun qisqartirish. Boshi muhimroq — oxiri kesiladi."""
    if len(text) <= max_chars:
        return text, False
    cut = text[:max_chars]
    nl = cut.rfind("\n")
    if nl > max_chars * 0.8:
        cut = cut[:nl]
    return cut + "\n\n[...matn uzunligi chegarasi sababli qisqartirildi...]", True
